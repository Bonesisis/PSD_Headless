"""Excel -> bounded editorial plan -> existing profile jobs. No renderer or AI dependency."""

from __future__ import annotations

import hashlib
import io
import json
import re
from datetime import date, datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import yaml
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, ConfigDict, Field

from .profiles import expand_profile

MAX_ROWS = 1000
ALIASES = {
    "plattform": "platform",
    "platform": "platform",
    "anzeigetitel": "title",
    "titel": "title",
    "title": "title",
    "datum": "date",
    "date": "date",
    "psd-datei": "template",
    "template": "template",
}


def iso_date(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ValueError("Datum benötigt TT.MM.JJJJ oder JJJJ-MM-TT")


def read_excel(raw: bytes) -> dict:
    """Read cells as data, never evaluate formulas, links or macros; bound ZIP expansion."""
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError("Excel-Datei darf maximal 5 MiB groß sein")
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(m.file_size for m in members) > 32 * 1024 * 1024:
                raise ValueError("Excel-Datei ist entpackt zu groß")
            if any("vbaproject" in m.filename.lower() for m in members):
                raise ValueError("Makrodateien werden nicht unterstützt; als XLSX speichern")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except (
        BadZipFile,
        KeyError,
        OSError,
        ParseError,
        DefusedXmlException,
        InvalidFileException,
    ) as exc:
        raise ValueError("Keine lesbare XLSX-Datei") from exc
    try:
        candidates = []
        for sheet in book.worksheets:
            for header_row, cells in enumerate(sheet.iter_rows(max_row=30, max_col=40), 1):
                columns = {}
                for cell in cells:
                    key = ALIASES.get(str(cell.value or "").strip().casefold())
                    if key:
                        if key in columns:
                            raise ValueError(f"Doppelte Spalte {key} in {sheet.title}")
                        columns[key] = cell.column
                if {"platform", "date"}.issubset(columns):
                    candidates.append((sheet, header_row, columns))
        if len(candidates) != 1:
            raise ValueError("Genau eine Tabelle mit Spalten Plattform und Datum wird benötigt")
        sheet, header, columns = candidates[0]
        # Do not trust a workbook's potentially incorrect dimension metadata.
        sheet.reset_dimensions()
        rows, errors = [], []
        for row, cells in enumerate(sheet.iter_rows(min_row=header + 1, max_col=40), header + 1):
            if row > header + MAX_ROWS:
                raise ValueError(f"Maximal {MAX_ROWS} Datenzeilen pro Upload")
            values = {k: cells[c - 1] for k, c in columns.items()}
            if not any(c.value is not None for c in cells):
                continue
            try:
                if any(c.data_type in {"f", "e"} for c in values.values()):
                    raise ValueError("Formeln in Auftragsfeldern bitte zuerst als Werte speichern")
                item = {k: str(c.value or "").strip() for k, c in values.items()}
                if any(len(v) > 2000 for v in item.values()):
                    raise ValueError("Auftragsfeld ist länger als 2000 Zeichen")
                item["date"] = iso_date(values["date"].value)
                if not item["platform"]:
                    raise ValueError("Plattform fehlt")
                rows.append({"row": row, **item})
            except ValueError as exc:
                errors.append({"row": row, "error": str(exc)})
        if not rows and not errors:
            raise ValueError("Excel enthält keine Aufträge")
        return {
            "sheet": sheet.title,
            "rows": rows,
            "errors": errors,
            "source_sha256": hashlib.sha256(raw).hexdigest(),
            "warnings": [
                "Bilddatei und Ausgabeordner aus Excel werden nicht verwendet. "
                "Ausgabeziele und Textebenen kommen aus den lokalen Profilen."
            ],
        }
    finally:
        book.close()


def catalog(root: str | Path) -> dict[str, dict]:
    """Only explicitly registered profiles in controller.json can be used by planners."""
    root = Path(root).resolve()
    path = root / "controller.json"
    if not path.is_file():
        raise ValueError("controller.json fehlt im Vorlagenordner; zuerst Profile registrieren")
    config = json.loads(path.read_text(encoding="utf-8-sig"))
    if config.get("version") != 1 or not isinstance(config.get("profiles"), dict):
        raise ValueError("Controller-Katalog benötigt version: 1 und profiles")
    result = {}
    for identifier, entry in config["profiles"].items():
        if not re.fullmatch(r"[a-z0-9_-]{1,64}", identifier):
            raise ValueError("Ungültige Profil-ID")
        profile_path = (root / entry["profile"]).resolve()
        if not profile_path.is_relative_to(root):
            raise ValueError("Profil verlässt den Vorlagenordner")
        profile = yaml.safe_load(profile_path.read_text(encoding="utf-8-sig"))
        result[identifier] = {
            **entry,
            "fields": list(profile["fields"]),
            "exports": len(profile.get("exports", [])),
        }
    if not result:
        raise ValueError("Keine Controller-Profile registriert")
    return result


class Assignment(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    row: int = Field(ge=1)
    profile_id: str
    fields: dict[str, str]


class Plan(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    assignments: list[Assignment] = Field(min_length=1, max_length=MAX_ROWS)


def planner_input(table: dict, profiles: dict) -> dict:
    return {
        "rows": table["rows"],
        "profiles": [
            {"id": k, "platform": v["platform"], "fields": v["fields"]} for k, v in profiles.items()
        ],
        "schema": Plan.model_json_schema(),
    }


def rule_plan(table: dict, profiles: dict) -> dict:
    assignments, errors = [], list(table["errors"])
    for row in table["rows"]:
        matches = [
            (k, v)
            for k, v in profiles.items()
            if str(v["platform"]).casefold() == row["platform"].casefold()
        ]
        if len(matches) != 1:
            errors.append({"row": row["row"], "error": "Plattform hat kein eindeutiges Profil"})
            continue
        identifier, profile = matches[0]
        fields = {k: row.get(k, "") for k in profile["fields"]}
        if any(not v for v in fields.values()):
            errors.append({"row": row["row"], "error": "Pflichtfeld des Profils fehlt"})
            continue
        assignments.append({"row": row["row"], "profile_id": identifier, "fields": fields})
    return {"assignments": assignments, "errors": errors}


def compile_plan(plan: dict, table: dict, profiles: dict, root: str | Path) -> list[dict]:
    """AI can supply fields and profile IDs, never commands, file paths or publication."""
    parsed = Plan.model_validate(plan)
    rows = {r["row"]: r for r in table["rows"]}
    if table["errors"] or sorted(a.row for a in parsed.assignments) != sorted(rows):
        raise ValueError(
            "Plan muss jede gültige Excel-Zeile genau einmal enthalten; Zeilenfehler beheben"
        )
    jobs = []
    for order_number, entry in enumerate(sorted(parsed.assignments, key=lambda a: a.row), 1):
        if entry.profile_id not in profiles:
            raise ValueError(f"Unbekanntes Profil: {entry.profile_id}")
        profile, row = profiles[entry.profile_id], rows[entry.row]
        if profile["platform"].casefold() != row["platform"].casefold():
            raise ValueError(f"Zeile {entry.row}: Profil passt nicht zur Plattform")
        if set(entry.fields) != set(profile["fields"]):
            raise ValueError(f"Zeile {entry.row}: Felder passen nicht zum Profil")
        for key, value in entry.fields.items():
            if not value.strip() or len(value) > 2000:
                raise ValueError(f"Zeile {entry.row}: {key} ist leer oder zu lang")
            if key == "date" and iso_date(value) != value:
                raise ValueError("Plan-Datum muss JJJJ-MM-TT sein")
        # Stable across retries, new content receives a new filename. Dates stay in profile outputs.
        identity = json.dumps([entry.profile_id, entry.row, entry.fields], sort_keys=True)
        suffix = hashlib.sha256(identity.encode()).hexdigest()[:12]
        title = re.sub(r"[^\w-]+", "-", row.get("title", "episode"), flags=re.UNICODE).strip("-_")
        name = f"{title[:55] or 'episode'}-{suffix}"
        job = {"profile": profile["profile"], "name": name, "fields": entry.fields}
        expanded = expand_profile(job, str(root), restricted=True)
        folder = Path(profile.get("output_folder", profile["platform"]))
        if folder.is_absolute() or len(folder.parts) != 1 or str(folder) in {".", ".."}:
            raise ValueError("Controller-Ausgabe benötigt einen festen Plattformordner")
        if any(Path(e["output"]).parent != folder for e in expanded):
            raise ValueError("Profil muss direkt in den registrierten Plattformordner exportieren")
        job["delivery"] = {
            "version": 1,
            "platform": str(folder),
            "order_number": order_number,
            "source_row": entry.row,
            "title": row.get("title", ""),
            "date_field": "date",
            "date": row["date"],
        }
        expand_profile(job, str(root), restricted=True)
        jobs.append(job)
    return jobs


def delivery_preview(jobs: list[dict], root: str | Path) -> list[dict]:
    result = []
    for job in jobs:
        exports = expand_profile(job, str(root), restricted=True)
        delivery = exports[0]["delivery"]
        result.append(
            {
                "row": delivery["row"],
                "title": delivery["title"],
                "folder": str(Path(exports[0]["output"]).parent).replace("\\", "/"),
                "files": [Path(e["output"]).name for e in exports],
            }
        )
    return result
